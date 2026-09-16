from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

_DEFAULT_RAW_DIR = Path("data/raw/onsv")
_DEFAULT_OUTPUT = Path("data/processed/onsv/lima_pedestrian_fatal_crashes.csv")
_CRASH_WORKBOOK = "fatal_crashes_2021_2025.xlsx"
_PEOPLE_WORKBOOK = "people_fatal_crashes_2021_2025.xlsx"
_HEADER_ROW = 5

_OUTPUT_COLUMNS = [
    "crash_id",
    "crash_date",
    "crash_time",
    "year",
    "month",
    "day",
    "district",
    "latitude",
    "longitude",
    "crash_class",
    "fatalities_total",
    "injured_total",
    "vehicles_damaged",
    "pedestrians_involved",
    "pedestrian_fatalities",
    "pedestrian_injured",
    "zone",
    "road_type",
    "road_network",
    "road_code",
    "weather_condition",
    "zoning",
    "road_characteristics",
    "road_longitudinal_profile",
    "road_surface",
    "vertical_sign_exists",
    "horizontal_sign_exists",
    "main_cause",
    "specific_cause",
]


def _key(value: Any) -> str:
    if value is None:
        return ""
    decomposed = unicodedata.normalize("NFKD", str(value))
    without_marks = "".join(character for character in decomposed if not unicodedata.combining(character))
    alphanumeric = re.sub(r"[^A-Za-z0-9]+", " ", without_marks.replace("_", " "))
    return " ".join(alphanumeric.upper().split())


def _text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).strip()


def _iter_table(path: Path, sheet_name: str) -> Iterator[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    header_values = next(
        worksheet.iter_rows(min_row=_HEADER_ROW, max_row=_HEADER_ROW, values_only=True)
    )
    headers = [_key(value) for value in header_values]
    try:
        for row in worksheet.iter_rows(min_row=_HEADER_ROW + 1, values_only=True):
            record = {header: value for header, value in zip(headers, row) if header}
            if _text(record.get("CODIGO SINIESTRO")):
                yield record
    finally:
        workbook.close()


def _is_lima_metropolitana(record: dict[str, Any]) -> bool:
    return _key(record.get("DEPARTAMENTO")) == "LIMA" and _key(record.get("PROVINCIA")) == "LIMA"


def _is_pedestrian(record: dict[str, Any]) -> bool:
    return _key(record.get("TIPO PERSONA")) == "PEATON"


def _iso_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _text(value)
    for date_format in ("%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, date_format).date().isoformat()
        except ValueError:
            continue
    return raw


def _iso_time(value: Any) -> str:
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0).isoformat()
    if isinstance(value, time):
        return value.replace(microsecond=0).isoformat()
    return _text(value)


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _integer(value: Any) -> int | None:
    number = _number(value)
    return None if number is None else int(number)


def _date_parts(date_value: str) -> tuple[str, str, str]:
    try:
        parsed = date.fromisoformat(date_value)
    except ValueError:
        return "", "", ""
    return str(parsed.year), f"{parsed.month:02d}", f"{parsed.day:02d}"


def build_lima_pedestrian_fatal_crashes(
    raw_dir: str | Path = _DEFAULT_RAW_DIR,
    output_path: str | Path = _DEFAULT_OUTPUT,
) -> dict[str, Any]:
    """Build one event-level row per fatal crash involving at least one pedestrian in Lima."""
    raw_root = Path(raw_dir)
    crash_path = raw_root / _CRASH_WORKBOOK
    people_path = raw_root / _PEOPLE_WORKBOOK
    if not crash_path.is_file():
        raise FileNotFoundError(crash_path)
    if not people_path.is_file():
        raise FileNotFoundError(people_path)

    pedestrian_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for person in _iter_table(people_path, "PERSONAS INVOLUCRADAS"):
        if not _is_lima_metropolitana(person) or not _is_pedestrian(person):
            continue
        crash_id = _text(person.get("CODIGO SINIESTRO"))
        severity = _key(person.get("GRAVEDAD"))
        pedestrian_counts[crash_id]["involved"] += 1
        if severity == "FALLECIDO":
            pedestrian_counts[crash_id]["fatalities"] += 1
        elif severity == "LESIONADO":
            pedestrian_counts[crash_id]["injured"] += 1

    records: list[dict[str, Any]] = []
    for crash in _iter_table(crash_path, "SINIESTROS"):
        if not _is_lima_metropolitana(crash):
            continue
        crash_id = _text(crash.get("CODIGO SINIESTRO"))
        pedestrians = pedestrian_counts.get(crash_id)
        if not pedestrians:
            continue

        crash_date = _iso_date(crash.get("FECHA SINIESTRO"))
        year, month, day = _date_parts(crash_date)
        records.append(
            {
                "crash_id": crash_id,
                "crash_date": crash_date,
                "crash_time": _iso_time(crash.get("HORA SINIESTRO")),
                "year": year,
                "month": month,
                "day": day,
                "district": _text(crash.get("DISTRITO")),
                "latitude": _number(crash.get("COORDENADAS LATITUD")),
                "longitude": _number(crash.get("COORDENADAS LONGITUD")),
                "crash_class": _text(crash.get("CLASE SINIESTRO")),
                "fatalities_total": _integer(crash.get("CANTIDAD DE FALLECIDOS")),
                "injured_total": _integer(crash.get("CANTIDAD DE LESIONADOS")),
                "vehicles_damaged": _integer(crash.get("CANTIDAD DE VEHICULOS DANADOS")),
                "pedestrians_involved": pedestrians["involved"],
                "pedestrian_fatalities": pedestrians["fatalities"],
                "pedestrian_injured": pedestrians["injured"],
                "zone": _text(crash.get("ZONA")),
                "road_type": _text(crash.get("TIPO DE VIA")),
                "road_network": _text(crash.get("RED VIAL")),
                "road_code": _text(crash.get("COD CARRETERA")),
                "weather_condition": _text(crash.get("CONDICION CLIMATICA")),
                "zoning": _text(crash.get("ZONIFICACION")),
                "road_characteristics": _text(crash.get("CARACTERISTICAS DE VIA")),
                "road_longitudinal_profile": _text(crash.get("PERFIL LONGITUDINAL VIA")),
                "road_surface": _text(crash.get("SUPERFICIE DE CALZADA")),
                "vertical_sign_exists": _text(crash.get("EXISTE SENAL VERTICAL")),
                "horizontal_sign_exists": _text(crash.get("EXISTE SENAL HORIZONTAL")),
                "main_cause": _text(crash.get("CAUSA FACTOR PRINCIPAL")),
                "specific_cause": _text(crash.get("CAUSA ESPECIFICA")),
            }
        )

    records.sort(key=lambda row: (row["crash_date"], row["crash_time"], row["crash_id"]))
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=_OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(records)

    valid_coordinates = sum(
        row["latitude"] is not None
        and row["longitude"] is not None
        and -90 <= row["latitude"] <= 90
        and -180 <= row["longitude"] <= 180
        for row in records
    )
    by_year = Counter(row["year"] for row in records if row["year"])
    return {
        "rows": len(records),
        "valid_coordinate_rows": valid_coordinates,
        "pedestrian_fatalities": sum(row["pedestrian_fatalities"] for row in records),
        "pedestrian_injured": sum(row["pedestrian_injured"] for row in records),
        "by_year": dict(sorted(by_year.items())),
        "output": str(destination),
    }


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Build the Lima pedestrian fatal-crash event dataset from ONSV sources."
    )
    parser.add_argument("--raw-dir", default=str(_DEFAULT_RAW_DIR))
    parser.add_argument("--output", default=str(_DEFAULT_OUTPUT))
    return parser


def main() -> None:
    args = _build_parser().parse_args()
    summary = build_lima_pedestrian_fatal_crashes(args.raw_dir, args.output)
    print(
        f"Wrote {summary['rows']} events to {summary['output']} "
        f"({summary['valid_coordinate_rows']} with valid coordinates)."
    )


if __name__ == "__main__":
    main()
