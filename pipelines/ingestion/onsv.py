from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from typing import Any, TypedDict
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

_REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
_DEFAULT_MANIFEST = _REPOSITORY_ROOT / "data" / "sources" / "onsv.json"
_DEFAULT_RAW_DIR = _REPOSITORY_ROOT / "data" / "raw" / "onsv"
_DEFAULT_REPORT = _REPOSITORY_ROOT / "data" / "interim" / "onsv" / "source-inventory.json"
_HEADER_SCAN_ROWS = 25


class SourceConfig(TypedDict):
    id: str
    role: str
    published: str
    filename: str
    url: str
    scope: str
    purpose: str


def load_sources(manifest_path: str | Path = _DEFAULT_MANIFEST) -> dict[str, SourceConfig]:
    """Load the ONSV source registry keyed by source id."""
    path = Path(manifest_path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    sources = payload.get("sources", [])
    return {source["id"]: source for source in sources}


def primary_source_ids(sources: dict[str, SourceConfig] | None = None) -> list[str]:
    registry = sources or load_sources()
    return [source_id for source_id, source in registry.items() if source["role"] == "primary"]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download_source(
    source: SourceConfig,
    output_dir: str | Path = _DEFAULT_RAW_DIR,
    *,
    force: bool = False,
) -> Path:
    """Download one registered source without placing partial files at the final path."""
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / source["filename"]
    if destination.exists() and not force:
        return destination

    partial = destination.with_suffix(destination.suffix + ".part")
    request = Request(
        source["url"],
        headers={"User-Agent": "Raster-Analytics-Web/1.0"},
    )
    with urlopen(request, timeout=120) as response, partial.open("wb") as target:  # noqa: S310
        shutil.copyfileobj(response, target)

    if not zipfile.is_zipfile(partial):
        partial.unlink(missing_ok=True)
        raise ValueError(f"Downloaded file for {source['id']} is not a valid XLSX workbook")

    partial.replace(destination)
    return destination


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_header(worksheet: Worksheet) -> tuple[int, list[str]]:
    best_row = 1
    best_values: list[str] = []
    best_score = -1
    max_row = min(worksheet.max_row or 1, _HEADER_SCAN_ROWS)

    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_row, values_only=True),
        start=1,
    ):
        values = [_display_value(value) for value in row]
        while values and not values[-1]:
            values.pop()
        score = sum(bool(value) for value in values)
        if score > best_score:
            best_row = row_index
            best_values = values
            best_score = score

    return best_row, best_values


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    """Return workbook structure and candidate headers without retaining record data."""
    workbook_path = Path(path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            header_row, columns = _find_header(worksheet)
            sheets.append(
                {
                    "name": worksheet.title,
                    "rows": max((worksheet.max_row or 0) - header_row, 0),
                    "columns": columns,
                    "header_row": header_row,
                }
            )
    finally:
        workbook.close()

    return {
        "filename": workbook_path.name,
        "size_bytes": workbook_path.stat().st_size,
        "sha256": _sha256(workbook_path),
        "sheets": sheets,
    }


def build_inventory(directory: str | Path, sources: dict[str, SourceConfig]) -> dict[str, Any]:
    root = Path(directory)
    entries: list[dict[str, Any]] = []
    for source_id, source in sources.items():
        path = root / source["filename"]
        if not path.exists():
            entries.append({"source_id": source_id, "status": "missing"})
            continue
        entries.append(
            {
                "source_id": source_id,
                "role": source["role"],
                "published": source["published"],
                "status": "available",
                **inspect_workbook(path),
            }
        )
    return {"sources": entries}


def _write_report(report: dict[str, Any], path: str | Path) -> Path:
    report_path = Path(path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return report_path


def sync_sources(
    *,
    include_all: bool = False,
    output_dir: str | Path = _DEFAULT_RAW_DIR,
    report_path: str | Path = _DEFAULT_REPORT,
    force: bool = False,
) -> dict[str, Any]:
    sources = load_sources()
    selected_ids = list(sources) if include_all else primary_source_ids(sources)
    selected = {source_id: sources[source_id] for source_id in selected_ids}

    for source in selected.values():
        download_source(source, output_dir, force=force)

    inventory = build_inventory(output_dir, selected)
    _write_report(inventory, report_path)
    return inventory


def _list_sources(sources: dict[str, SourceConfig]) -> None:
    for source_id, source in sources.items():
        print(f"{source_id}\t{source['role']}\t{source['published']}\t{source['url']}")


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Download and inspect registered ONSV datasets.")
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("list", help="List registered ONSV data sources.")

    sync_parser = subparsers.add_parser(
        "sync",
        help="Download sources and write an XLSX inventory.",
    )
    sync_parser.add_argument(
        "--all",
        action="store_true",
        help="Include supporting/context sources.",
    )
    sync_parser.add_argument(
        "--force",
        action="store_true",
        help="Replace existing downloads.",
    )
    sync_parser.add_argument("--output-dir", default=str(_DEFAULT_RAW_DIR))
    sync_parser.add_argument("--report", default=str(_DEFAULT_REPORT))

    inspect_parser = subparsers.add_parser("inspect", help="Inspect already downloaded workbooks.")
    inspect_parser.add_argument("--input-dir", default=str(_DEFAULT_RAW_DIR))
    inspect_parser.add_argument("--report", default=str(_DEFAULT_REPORT))
    inspect_parser.add_argument(
        "--all",
        action="store_true",
        help="Include supporting/context sources.",
    )
    return parser


def main() -> None:
    parser = _build_parser()
    args = parser.parse_args()
    sources = load_sources()

    if args.command == "list":
        _list_sources(sources)
        return

    selected_ids = list(sources) if args.all else primary_source_ids(sources)
    selected = {source_id: sources[source_id] for source_id in selected_ids}

    if args.command == "sync":
        for source in selected.values():
            download_source(source, args.output_dir, force=args.force)
        report = build_inventory(args.output_dir, selected)
    else:
        report = build_inventory(args.input_dir, selected)

    report_path = _write_report(report, args.report)
    print(f"Inventory written to {report_path}")


if __name__ == "__main__":
    main()
